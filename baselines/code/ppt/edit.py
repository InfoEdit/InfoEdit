"""Code-based PPTX infographic editing.

Sister script to ``code/html/edit.py``: instead of HTML, we operate on
PowerPoint slides via python-pptx.

Pipeline per item:
  1. Look up the slide index in the master deck via ``id_mapping.csv`` and
     dump that slide's structure (shapes, text, positions in EMU) as text.
  2. Send the dump (+ optional rendered PNG in ``code_image`` mode) and the
     edit instruction to Gemini, which returns a ``def edit(prs):`` function.
  3. Build a single-slide working deck from the master, ``exec`` the LLM code,
     call ``edit(prs)``, save the modified ``.pptx``.
  4. Render the modified ``.pptx`` to PNG via LibreOffice (headless) + pdf2image.

Downstream evaluation reads ``{id}_edited_1.png`` exactly as the HTML path does.
"""

import argparse
import csv
import io
import json
import multiprocessing as mp
import os
import random
import subprocess
import time
import types as pytypes
from concurrent.futures import ProcessPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path

from PIL import Image
try:                      # Vertex-only; unused on the openai-proxy branch
    from google import genai
    from google.genai import types
except ImportError:       # pragma: no cover
    genai = types = None
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Emu, Inches, Pt
from tqdm import tqdm
try:
    import wandb
except ImportError:   # optional: only needed when W&B logging is enabled
    wandb = None

# Reach llm_client.py at the repo root.
import sys as _sys, pathlib as _pl
_sys.path.insert(0, str(_pl.Path(__file__).resolve().parents[3]))
from llm_client import chat_text


DEFAULT_MODEL_PATH = "gemini-3.1-pro-preview"
NUM_VARIANTS = 1
OPERATIONS = ("add", "delete", "swap_intra", "swap_inter", "text_expand", "aspect_ratio")
INPUT_MODES = ("code", "code_image")
DEFAULT_MASTER_NAME = "tempates_gallery_0430.pptx"
DEFAULT_MAPPING_NAME = "id_mapping.csv"
DEFAULT_RENDER_DPI = 100  # matches the 1334x750 PNGs already shipped for v7

# Same per-orientation AR pool as the HTML path. python-pptx is unconstrained
# by any model's supported-AR list, so we just pick a representative AR.
PPT_AR_BY_ORIENTATION = {
    "landscape": ["3:2", "16:9", "4:3", "21:9"],
    "portrait":  ["2:3", "9:16", "3:4"],
    "square":    ["1:1"],
}


def ppt_ar_for_orientation(orientation: str, idx: int):
    options = PPT_AR_BY_ORIENTATION.get(orientation)
    if not options:
        return None
    return options[idx % len(options)]


def _guess_mime(path):
    lower = path.lower()
    if lower.endswith(".jpg") or lower.endswith(".jpeg"):
        return "image/jpeg"
    if lower.endswith(".webp"):
        return "image/webp"
    if lower.endswith(".gif"):
        return "image/gif"
    return "image/png"


def _thinking_config(model_path: str):
    if model_path.startswith("gemini-3"):
        return types.ThinkingConfig(thinking_level="High", include_thoughts=True)
    return types.ThinkingConfig(thinking_budget=-1, include_thoughts=True)


def _thinking_config_json(model_path: str) -> dict:
    if model_path.startswith("gemini-3"):
        return {"thinkingLevel": "High", "includeThoughts": True}
    return {"thinkingBudget": -1, "includeThoughts": True}


EDIT_SYSTEM_INSTRUCTION = (
    "You are an expert Python engineer who edits PowerPoint infographic slides via python-pptx. "
    "You will receive a structured DUMP of a single-slide deck and an editing instruction. "
    "Return ONLY a Python module that defines a single function with this exact signature:\n\n"
    "    def edit(prs):  # prs is a python-pptx Presentation containing exactly one slide (prs.slides[0])\n"
    "        ...\n\n"
    "Hard requirements:\n"
    "  - Output MUST be raw Python source. No prose, no explanations, no markdown code fences.\n"
    "  - Use ONLY python-pptx APIs. The following names are pre-imported in your namespace, do not re-import: "
    "    Presentation, Inches, Pt, Emu, Cm, RGBColor, MSO_SHAPE, MSO_SHAPE_TYPE, PP_ALIGN.\n"
    "  - The deck has exactly one slide accessible as prs.slides[0]. Reference shapes by iterating "
    "    prs.slides[0].shapes (recurse into groups via shape.shape_type == MSO_SHAPE_TYPE.GROUP) "
    "    and matching shape_id from the dump.\n"
    "  - Preserve every shape and property unrelated to the requested edit.\n"
    "  - All positions and sizes are in EMU (914400 EMU = 1 inch). Use Emu(int) or Inches(float) explicitly.\n"
    "  - When TARGET_DIMENSIONS_EMU is provided (aspect-ratio edits), the working deck's slide_width / "
    "    slide_height have ALREADY been set to those dimensions before edit() is called. Your job is to "
    "    rescale and reposition every shape (recursively into groups) so the layout fills the new canvas "
    "    without distortion or overflow. Do NOT change prs.slide_width / prs.slide_height yourself.\n"
    "  - Otherwise do NOT touch slide dimensions.\n"
    "  - Do NOT call prs.save(); the caller handles persistence.\n"
    "  - Do NOT print, log, or raise; on any unrecoverable problem, simply leave the slide unchanged.\n"
    "  - If you need lower-level OXML helpers, the correct import paths are: "
    "`from pptx.oxml.xmlchemy import OxmlElement` and `from pptx.oxml.ns import qn`. "
    "Do NOT use `from pptx.oxml import OxmlElement` — that path does not exist."
)


def input_file_for(base, operation):
    stem, ext = os.path.splitext(base)
    return f"{stem}.{operation}{ext}" if ext else f"{base}.{operation}"


def output_dir_for(base, operation):
    return os.path.join(base, operation)


def meta_path_for(image_path):
    return os.path.splitext(image_path)[0] + ".meta.json"


def load_id_mapping(csv_path):
    """Read id_mapping.csv → dict[new_id:int] = old_id:int (1-based slide index)."""
    mapping = {}
    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            try:
                new_id = int(row["new_id"])
                old_id = int(row["old_id"])
            except (KeyError, ValueError):
                continue
            mapping[new_id] = old_id
    return mapping


def load_meta(image_path):
    p = meta_path_for(image_path)
    if not os.path.exists(p):
        return None
    try:
        with open(p, "r", encoding="utf-8") as f:
            return json.load(f)
    except (OSError, json.JSONDecodeError):
        return None


def parse_aspect_ratio(target_ar):
    if not target_ar or ":" not in target_ar:
        return None
    try:
        a, b = target_ar.split(":")
        a, b = float(a), float(b)
    except ValueError:
        return None
    if a <= 0 or b <= 0:
        return None
    return a, b


def compute_new_dims_emu(orig_w_emu, orig_h_emu, target_ar):
    """Pick new (w, h) in EMU for an aspect-ratio edit by preserving total area."""
    parsed = parse_aspect_ratio(target_ar)
    if not parsed:
        return orig_w_emu, orig_h_emu
    a, b = parsed
    area = orig_w_emu * orig_h_emu
    new_w = int(round((area * a / b) ** 0.5))
    new_h = int(round(new_w * b / a))
    return new_w, new_h


def load_prompts(file_path, prompt_index=0):
    """Mirror of code/html/edit.py: load_prompts; emits target_aspect_ratio."""
    data = []
    with open(file_path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            item = json.loads(line)
            if not item.get("image_path"):
                continue

            prompts_list = item.get("generated_edit_prompts")
            if isinstance(prompts_list, list) and prompts_list:
                idx = prompt_index if prompt_index < len(prompts_list) else 0
                item["generated_edit_prompt"] = str(prompts_list[idx]).strip()
                target_orientations = item.get("target_orientations")
                if isinstance(target_orientations, list) and len(target_orientations) > idx:
                    orientation = str(target_orientations[idx]).strip()
                    ar = ppt_ar_for_orientation(orientation, idx)
                    if ar:
                        item["target_aspect_ratio"] = ar
            elif item.get("generated_edit_prompt"):
                item["generated_edit_prompt"] = str(item["generated_edit_prompt"]).strip()
            else:
                continue

            if not item["generated_edit_prompt"]:
                continue
            data.append(item)
    return data


# ---------------------------------------------------------------------------
# Slide structure dump (the LLM's "source code")
# ---------------------------------------------------------------------------

def _shape_kind_name(sh):
    try:
        return str(sh.shape_type).rsplit(".", 1)[-1]
    except Exception:
        return "UNKNOWN"


def _safe_text_lines(sh, indent_pad):
    out = []
    if not sh.has_text_frame:
        return out
    tf = sh.text_frame
    for j, para in enumerate(tf.paragraphs):
        runs_text = "".join((r.text or "") for r in para.runs) or (para.text or "")
        runs_text = runs_text.replace("\n", " ")
        if not runs_text:
            continue
        # Surface dominant font of first run for size cues.
        first_run = para.runs[0] if para.runs else None
        font_bits = []
        if first_run is not None and first_run.font is not None:
            f = first_run.font
            try:
                if f.name: font_bits.append(f"name={f.name}")
            except Exception: pass
            try:
                if f.size: font_bits.append(f"size={f.size.pt}pt")
            except Exception: pass
            try:
                if f.bold is not None: font_bits.append(f"bold={f.bold}")
            except Exception: pass
        font_str = (" font=(" + ", ".join(font_bits) + ")") if font_bits else ""
        out.append(f"{indent_pad}  text[p{j}]: {runs_text!r}{font_str}")
    return out


def _dump_shape(sh, indent=0, lines=None):
    if lines is None:
        lines = []
    pad = " " * indent
    kind = _shape_kind_name(sh)
    try:
        sid = sh.shape_id
    except Exception:
        sid = "?"
    try:
        name = sh.name
    except Exception:
        name = "?"
    lines.append(f"{pad}- shape_id={sid} type={kind} name={name!r}")
    try:
        lines.append(
            f"{pad}  pos=(L={sh.left}, T={sh.top}) size=(W={sh.width}, H={sh.height})"
        )
    except Exception:
        pass
    lines.extend(_safe_text_lines(sh, pad))
    # Picture: just note presence (no need for image bytes in dump).
    try:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            lines.append(f"{pad}  picture: <embedded image>")
    except Exception:
        pass
    # Group: recurse.
    try:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            lines.append(f"{pad}  group_children:")
            for child in sh.shapes:
                _dump_shape(child, indent=indent + 4, lines=lines)
    except Exception:
        pass
    return lines


def dump_slide(slide, slide_w_emu, slide_h_emu):
    lines = [
        f"SLIDE_SIZE_EMU: width={slide_w_emu}, height={slide_h_emu}  "
        f"# 914400 EMU = 1 inch ({slide_w_emu/914400:.2f}in x {slide_h_emu/914400:.2f}in)",
        f"SLIDE_LAYOUT: {slide.slide_layout.name!r}",
        f"TOP_LEVEL_SHAPE_COUNT: {len(slide.shapes)}",
        "",
        "SHAPES:",
    ]
    for sh in slide.shapes:
        _dump_shape(sh, indent=2, lines=lines)
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Prompt construction + LLM output cleanup
# ---------------------------------------------------------------------------

def build_user_prompt(
    slide_dump, edit_prompt, target_w_emu=None, target_h_emu=None, with_image=False,
):
    parts = [f"EDITING_INSTRUCTION:\n{edit_prompt}"]
    if target_w_emu and target_h_emu:
        parts.append(
            f"TARGET_DIMENSIONS_EMU: width={target_w_emu}, height={target_h_emu}  "
            f"# already applied to prs.slide_width/height before edit() runs"
        )
    if with_image:
        parts.append(
            "RENDERED_PREVIEW: A PNG rendered from the ORIGINAL deck is attached for visual context "
            "(layout, colors, positions). All modifications MUST be expressed in the returned Python "
            "function — do not output an image."
        )
    parts.append(f"SLIDE_DUMP:\n{slide_dump}")
    return "\n\n".join(parts)


_IMPORT_FIXUPS = (
    # OxmlElement lives in pptx.oxml.xmlchemy, not pptx.oxml.
    ("from pptx.oxml import OxmlElement", "from pptx.oxml.xmlchemy import OxmlElement"),
)


def clean_python_output(text):
    """Strip stray markdown fences in case the model ignores the no-fence rule."""
    if not text:
        return ""
    text = text.strip()
    if text.startswith("```"):
        # Drop the opening fence line (e.g. ``` or ```python).
        nl = text.find("\n")
        if nl != -1:
            text = text[nl + 1 :]
        if text.rstrip().endswith("```"):
            text = text.rstrip()[:-3]
    text = text.strip()
    for bad, good in _IMPORT_FIXUPS:
        if bad in text:
            text = text.replace(bad, good)
    return text


# ---------------------------------------------------------------------------
# Single-slide deck building + edit application
# ---------------------------------------------------------------------------

def _drop_other_slides(prs, keep_idx_0based):
    """Mutate prs in place, keeping only the slide at keep_idx_0based."""
    sldIdLst = prs.slides._sldIdLst  # CT_SlideIdList (lxml element)
    sldIds = list(sldIdLst)
    if keep_idx_0based < 0 or keep_idx_0based >= len(sldIds):
        raise IndexError(
            f"slide index {keep_idx_0based} out of range ({len(sldIds)} slides in master)"
        )
    for i, e in enumerate(sldIds):
        if i == keep_idx_0based:
            continue
        try:
            prs.part.drop_rel(e.rId)
        except Exception:
            pass
        sldIdLst.remove(e)


def _exec_edit_code(edit_code):
    """Compile + exec LLM code, returning the edit() callable.

    Pre-imports python-pptx names so the model doesn't need ``import`` statements.
    """
    module = pytypes.ModuleType("llm_edit")
    module.__dict__.update({
        "Presentation": Presentation,
        "Inches": Inches, "Pt": Pt, "Emu": Emu, "Cm": Cm,
        "RGBColor": RGBColor,
        "MSO_SHAPE": MSO_SHAPE,
        "MSO_SHAPE_TYPE": MSO_SHAPE_TYPE,
        "PP_ALIGN": PP_ALIGN,
    })
    exec(compile(edit_code, "<llm_edit>", "exec"), module.__dict__)
    fn = module.__dict__.get("edit")
    if not callable(fn):
        raise ValueError("LLM output did not define a callable edit(prs)")
    return fn


def apply_edit_and_save(
    master_bytes,
    slide_idx_0based,
    edit_code,
    output_pptx_path,
    target_w_emu=None,
    target_h_emu=None,
):
    """Build a single-slide deck from master_bytes, run LLM edit, save PPTX.

    Returns (ok, error_msg). On success, ok=True and error_msg=None.
    """
    try:
        bio = io.BytesIO(master_bytes)
        prs = Presentation(bio)
        _drop_other_slides(prs, slide_idx_0based)

        if target_w_emu and target_h_emu:
            prs.slide_width = int(target_w_emu)
            prs.slide_height = int(target_h_emu)

        edit_fn = _exec_edit_code(edit_code)
        edit_fn(prs)
        prs.save(output_pptx_path)
        return True, None
    except Exception as e:
        return False, f"{type(e).__name__}: {e}"


# ---------------------------------------------------------------------------
# Output paths + skip-detection
# ---------------------------------------------------------------------------

def edited_paths(output_dir, item_id, var_idx=1):
    base = os.path.join(output_dir, f"{item_id}_edited_{var_idx}")
    return (
        base + ".pptx",
        base + ".png",
        base + "_thoughts.txt",
        base + ".render.json",
        base + ".edit.py",   # also dump the LLM code for debugging
    )


def all_code_exist(output_dir, item_id):
    """Primary success indicator: every variant has a non-empty .edit.py file.

    pptx and png are downstream/consequential outputs derived from code, so we
    track LLM-phase progress by code presence only.
    """
    for v in range(1, NUM_VARIANTS + 1):
        code_path = edited_paths(output_dir, item_id, v)[4]
        if not os.path.exists(code_path):
            return False
        try:
            if os.path.getsize(code_path) == 0:
                return False
        except OSError:
            return False
    return True


def all_pngs_exist(output_dir, item_id):
    return all(
        os.path.exists(edited_paths(output_dir, item_id, v)[1])
        for v in range(1, NUM_VARIANTS + 1)
    )


def all_pptx_exist(output_dir, item_id):
    return all(
        os.path.exists(edited_paths(output_dir, item_id, v)[0])
        for v in range(1, NUM_VARIANTS + 1)
    )


# ---------------------------------------------------------------------------
# Online (synchronous) editing
# ---------------------------------------------------------------------------

def _read_master_bytes(master_path):
    with open(master_path, "rb") as f:
        return f.read()


def process_edits(
    clients, model_path, data_list, output_dir, master_pptx, id_mapping,
    worker_id=0, input_mode="code",
):
    os.makedirs(output_dir, exist_ok=True)
    max_retries = 3
    success_count = 0
    fail_count = 0
    with_image = input_mode == "code_image"

    master_bytes = _read_master_bytes(master_pptx)
    # We need a Presentation object to dump slides. Loading once here is enough
    # because we never mutate this instance — apply_edit_and_save reloads from
    # master_bytes each time.
    master_prs_for_dump = Presentation(io.BytesIO(master_bytes))
    master_slides = list(master_prs_for_dump.slides)
    slide_w = master_prs_for_dump.slide_width
    slide_h = master_prs_for_dump.slide_height

    desc = f"Worker-{worker_id}" if worker_id else "Editing PPT"
    for item in tqdm(data_list, desc=desc, position=worker_id):
        item_id = str(item.get("id"))
        image_path = item.get("image_path")
        edit_prompt = item.get("generated_edit_prompt")
        target_ar = item.get("target_aspect_ratio")

        try:
            new_id_int = int(item_id)
        except ValueError:
            tqdm.write(f"[Warning] non-integer id {item_id}, skipping")
            fail_count += 1
            continue
        old_id = id_mapping.get(new_id_int)
        if old_id is None:
            tqdm.write(f"[Warning] id {item_id} not found in id_mapping, skipping")
            fail_count += 1
            continue
        slide_idx = old_id - 1
        if not (0 <= slide_idx < len(master_slides)):
            tqdm.write(f"[Warning] slide index {slide_idx} out of range for id {item_id}")
            fail_count += 1
            continue

        slide = master_slides[slide_idx]
        slide_dump = dump_slide(slide, slide_w, slide_h)

        source_image = None
        if with_image:
            if not image_path or not os.path.exists(image_path):
                tqdm.write(f"[Warning] Rendered preview not found for ID {item_id}: {image_path}")
                fail_count += 1
                continue
            try:
                source_image = Image.open(image_path)
            except Exception as e:
                tqdm.write(f"[Error] Failed to open preview for ID {item_id}: {e}")
                fail_count += 1
                continue

        new_w_emu, new_h_emu = (
            compute_new_dims_emu(slide_w, slide_h, target_ar) if target_ar else (slide_w, slide_h)
        )

        user_prompt = build_user_prompt(
            slide_dump, edit_prompt,
            target_w_emu=new_w_emu if target_ar else None,
            target_h_emu=new_h_emu if target_ar else None,
            with_image=with_image,
        )

        item_variant_success = 0
        for var_idx in range(1, NUM_VARIANTS + 1):
            pptx_save, png_save, txt_save, render_save, code_save = edited_paths(
                output_dir, item_id, var_idx
            )

            # Phase 1 — ensure the LLM code exists on disk. This is the primary
            # success criterion. If a prior run already saved code we reuse it
            # without calling the LLM; otherwise we call Gemini and persist the
            # response. Apply (Phase 2) is treated as a downstream side-effect
            # and never gates success.
            edited_code = ""
            thoughts_to_save = []

            if os.path.exists(code_save):
                try:
                    with open(code_save, "r", encoding="utf-8") as f:
                        edited_code = f.read()
                except OSError:
                    edited_code = ""

            if not edited_code.strip():
                edited_code = ""
                for attempt in range(max_retries):
                    try:
                        resp_code = chat_text(
                            model_path,
                            user_prompt,
                            images=[source_image] if source_image is not None else None,
                            system=EDIT_SYSTEM_INSTRUCTION,
                        )
                        # The proxy exposes no separate reasoning stream.
                        resp_thoughts = []

                        resp_code = clean_python_output(resp_code)
                        if resp_code:
                            edited_code = resp_code
                            thoughts_to_save = resp_thoughts
                            break

                        tqdm.write(
                            f"[Warning] Empty code for ID {item_id} (Variant {var_idx}), "
                            f"retrying ({attempt + 1}/{max_retries})..."
                        )
                        time.sleep(2)
                    except Exception as e:
                        tqdm.write(
                            f"[Error] ID {item_id} (Variant {var_idx}, Attempt {attempt + 1}/{max_retries}): {e}"
                        )
                        time.sleep(2)

                if not edited_code:
                    # All retries exhausted; no code → variant fails.
                    time.sleep(1)
                    continue

                with open(code_save, "w", encoding="utf-8") as f:
                    f.write(edited_code)
                if thoughts_to_save:
                    with open(txt_save, "w", encoding="utf-8") as f:
                        f.write("\n\n".join(thoughts_to_save))

            # Code is now on disk → variant counts as success regardless of
            # what happens in Phase 2.
            item_variant_success += 1

            # Phase 2 — apply the code to produce the .pptx. Best-effort: a
            # failure here is logged but does NOT trigger an LLM retry and
            # does NOT mark the variant as failed.
            if not os.path.exists(pptx_save):
                ok, err = apply_edit_and_save(
                    master_bytes=master_bytes,
                    slide_idx_0based=slide_idx,
                    edit_code=edited_code,
                    output_pptx_path=pptx_save,
                    target_w_emu=new_w_emu if target_ar else None,
                    target_h_emu=new_h_emu if target_ar else None,
                )
                if ok:
                    with open(render_save, "w", encoding="utf-8") as f:
                        json.dump({"width_emu": new_w_emu, "height_emu": new_h_emu}, f)
                else:
                    tqdm.write(
                        f"[Warning] apply_edit failed for ID {item_id} (Variant {var_idx}): "
                        f"{err} (code preserved)"
                    )
                    try:
                        if os.path.exists(pptx_save):
                            os.remove(pptx_save)
                    except OSError:
                        pass
            time.sleep(1)

        if item_variant_success == NUM_VARIANTS:
            success_count += 1
        else:
            fail_count += 1

    return success_count, fail_count


def _worker_entry(
    worker_id, api_keys, model_path, data_shard, output_dir,
    master_pptx, id_mapping, result_queue, input_mode,
):
    try:
        clients = [None]  # unused: llm_client holds the proxy client
        success, fail = process_edits(
            clients=clients,
            model_path=model_path,
            data_list=data_shard,
            output_dir=output_dir,
            master_pptx=master_pptx,
            id_mapping=id_mapping,
            worker_id=worker_id,
            input_mode=input_mode,
        )
        result_queue.put((worker_id, success, fail))
    except Exception as e:
        print(f"[Worker-{worker_id}] Fatal error: {e}")
        result_queue.put((worker_id, 0, len(data_shard)))


def run_parallel(
    num_workers, api_keys, model_path, data, output_dir,
    master_pptx, id_mapping, input_mode="code",
):
    shards = [[] for _ in range(num_workers)]
    for i, item in enumerate(data):
        shards[i % num_workers].append(item)

    ctx = mp.get_context("spawn")
    result_queue = ctx.Queue()
    processes = []

    for wid in range(num_workers):
        if not shards[wid]:
            continue
        p = ctx.Process(
            target=_worker_entry,
            args=(
                wid, api_keys, model_path, shards[wid], output_dir,
                master_pptx, id_mapping, result_queue, input_mode,
            ),
        )
        p.start()
        processes.append(p)

    total_success = 0
    total_fail = 0
    for _ in processes:
        wid, s, f = result_queue.get()
        total_success += s
        total_fail += f

    for p in processes:
        p.join()

    return total_success, total_fail


# ---------------------------------------------------------------------------
# Vertex AI batch editing
# ---------------------------------------------------------------------------

def _batch_state_name(state):
    return getattr(state, "name", None) or str(state)


def process_edits_batch(
    project_id, location, bucket_uri, model_path, data_list, output_dir,
    master_pptx, id_mapping, poll_interval=30, input_mode="code",
):
    """Submit all PPT-editing requests as a single Vertex AI Gemini batch job."""
    import fsspec  # lazy — gcsfs only needed in batch mode
    from google.genai.types import CreateBatchJobConfig

    os.makedirs(output_dir, exist_ok=True)
    bucket_uri = bucket_uri.rstrip("/")
    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    with_image = input_mode == "code_image"

    master_bytes = _read_master_bytes(master_pptx)
    master_prs_for_dump = Presentation(io.BytesIO(master_bytes))
    master_slides = list(master_prs_for_dump.slides)
    slide_w = master_prs_for_dump.slide_width
    slide_h = master_prs_for_dump.slide_height

    client = genai.Client(vertexai=True, project=project_id, location=location)
    fs = fsspec.filesystem("gcs")

    input_jsonl_uri = f"{bucket_uri}/ppt_batch_input_{timestamp}.jsonl"
    output_prefix = f"{bucket_uri}/ppt_batch_output_{timestamp}"
    images_gcs_dir = f"{bucket_uri}/ppt_batch_images_{timestamp}" if with_image else None

    # full user prompt -> (item_id, slide_idx_0based, new_w_emu, new_h_emu, target_ar)
    prompt_to_item = {}
    cached_success = 0  # items where code already exists; LLM skipped

    print(f"[Batch] Writing requests to {input_jsonl_uri} (input_mode={input_mode})")
    written = 0
    with fs.open(input_jsonl_uri, "w", encoding="utf-8") as f_out:
        for item in data_list:
            item_id = str(item.get("id"))
            image_path = item.get("image_path")
            edit_prompt = item.get("generated_edit_prompt")
            target_ar = item.get("target_aspect_ratio")

            try:
                new_id_int = int(item_id)
            except ValueError:
                print(f"[Batch] non-integer id {item_id}, skipping")
                continue
            old_id = id_mapping.get(new_id_int)
            if old_id is None:
                print(f"[Batch] id {item_id} not found in id_mapping, skipping")
                continue
            slide_idx = old_id - 1
            if not (0 <= slide_idx < len(master_slides)):
                print(f"[Batch] slide index {slide_idx} out of range for id {item_id}, skipping")
                continue

            new_w_emu, new_h_emu = (
                compute_new_dims_emu(slide_w, slide_h, target_ar) if target_ar else (slide_w, slide_h)
            )

            # Skip the LLM call if a prior run already saved the edit code.
            # Code presence is the success criterion — apply is best-effort and
            # never affects the cached_success count.
            pptx_save, _, _, render_save, code_save = edited_paths(output_dir, item_id, 1)
            cached_code = ""
            if os.path.exists(code_save):
                try:
                    with open(code_save, "r", encoding="utf-8") as fc:
                        cached_code = fc.read()
                except OSError:
                    cached_code = ""
            if cached_code.strip():
                cached_success += 1
                if not os.path.exists(pptx_save):
                    ok, err = apply_edit_and_save(
                        master_bytes=master_bytes,
                        slide_idx_0based=slide_idx,
                        edit_code=cached_code,
                        output_pptx_path=pptx_save,
                        target_w_emu=new_w_emu if target_ar else None,
                        target_h_emu=new_h_emu if target_ar else None,
                    )
                    if ok:
                        with open(render_save, "w", encoding="utf-8") as fr:
                            json.dump({"width_emu": new_w_emu, "height_emu": new_h_emu}, fr)
                        print(f"[Batch] ID {item_id}: applied from cached code, skipping LLM")
                    else:
                        print(f"[Batch] ID {item_id}: cached code failed to apply ({err}); code preserved, LLM skipped")
                        try:
                            if os.path.exists(pptx_save):
                                os.remove(pptx_save)
                        except OSError:
                            pass
                else:
                    print(f"[Batch] ID {item_id}: code and pptx already present, skipping LLM")
                continue

            slide_dump = dump_slide(master_slides[slide_idx], slide_w, slide_h)

            image_uri = None
            if with_image:
                if not image_path:
                    print(f"[Batch] Skipping ID {item_id}: missing image_path")
                    continue
                if image_path.startswith("gs://"):
                    image_uri = image_path
                else:
                    if not os.path.exists(image_path):
                        print(f"[Batch] Skipping ID {item_id}: rendered preview not found: {image_path}")
                        continue
                    ext = os.path.splitext(image_path)[1] or ".png"
                    image_uri = f"{images_gcs_dir}/{item_id}{ext}"
                    try:
                        with open(image_path, "rb") as src, fs.open(image_uri, "wb") as dst:
                            dst.write(src.read())
                    except Exception as e:
                        print(f"[Batch] Failed to upload preview for ID {item_id}: {e}")
                        continue

            user_prompt = build_user_prompt(
                slide_dump, edit_prompt,
                target_w_emu=new_w_emu if target_ar else None,
                target_h_emu=new_h_emu if target_ar else None,
                with_image=with_image,
            )
            prompt_to_item[user_prompt] = (item_id, slide_idx, new_w_emu, new_h_emu, target_ar)

            user_parts = [{"text": user_prompt}]
            if with_image and image_uri:
                user_parts.append({
                    "file_data": {"file_uri": image_uri, "mime_type": _guess_mime(image_path)}
                })

            payload = {
                "request": {
                    "systemInstruction": {"parts": [{"text": EDIT_SYSTEM_INSTRUCTION}]},
                    "contents": [{"role": "user", "parts": user_parts}],
                    "generationConfig": {
                        "thinkingConfig": _thinking_config_json(model_path),
                    },
                }
            }
            f_out.write(json.dumps(payload, ensure_ascii=False) + "\n")
            written += 1

    if cached_success:
        print(f"[Batch] Reused {cached_success} cached code files (LLM skipped).")

    if written == 0:
        print("[Batch] No new requests to submit (all handled from cache or skipped).")
        return cached_success, len(data_list) - cached_success

    print(f"[Batch] Submitting batch job (model={model_path}, dest={output_prefix})")
    batch_job = client.batches.create(
        model=model_path,
        src=input_jsonl_uri,
        config=CreateBatchJobConfig(dest=output_prefix),
    )
    print(f"[Batch] Job name: {batch_job.name}")

    running = {"JOB_STATE_RUNNING", "JOB_STATE_PENDING", "JOB_STATE_QUEUED"}
    while _batch_state_name(batch_job.state) in running:
        print(f"[Batch] state={_batch_state_name(batch_job.state)}; sleeping {poll_interval}s")
        time.sleep(poll_interval)
        batch_job = client.batches.get(name=batch_job.name)

    final_state = _batch_state_name(batch_job.state)
    if final_state != "JOB_STATE_SUCCEEDED":
        print(f"[Batch] Job {batch_job.name} failed: state={final_state} error={batch_job.error}")
        return cached_success, len(data_list) - cached_success

    dest_uri = batch_job.dest.gcs_uri
    pred_paths = fs.glob(f"{dest_uri}/*/predictions.jsonl")
    if not pred_paths:
        print(f"[Batch] No predictions.jsonl under {dest_uri}")
        return cached_success, len(data_list) - cached_success

    success = 0
    fail = 0
    seen = set()
    for pred_path in pred_paths:
        with fs.open(f"gs://{pred_path}", "r", encoding="utf-8") as f_in:
            for line in f_in:
                line = line.strip()
                if not line:
                    continue
                try:
                    row = json.loads(line)
                except json.JSONDecodeError as e:
                    print(f"[Batch] Skipping malformed line: {e}")
                    fail += 1
                    continue

                prompt_text = None
                try:
                    parts = row["request"]["contents"][0]["parts"]
                    for p in parts:
                        if p.get("text"):
                            prompt_text = p["text"]
                except (KeyError, IndexError, TypeError):
                    pass

                matched = prompt_to_item.get(prompt_text) if prompt_text else None
                if not matched:
                    print("[Batch] Could not match response back to any input item")
                    fail += 1
                    continue
                item_id, slide_idx, new_w_emu, new_h_emu, target_ar = matched
                seen.add(prompt_text)

                response = row.get("response") or {}
                edited_code = ""
                thoughts = []
                for cand in response.get("candidates", []) or []:
                    content = cand.get("content") or {}
                    for part in content.get("parts", []) or []:
                        text = part.get("text")
                        if not text:
                            continue
                        if part.get("thought"):
                            thoughts.append(text)
                        else:
                            edited_code += text

                edited_code = clean_python_output(edited_code)
                if not edited_code:
                    print(f"[Batch][Failed] ID {item_id}: empty code in response")
                    fail += 1
                    continue

                pptx_save, _, txt_save, render_save, code_save = edited_paths(
                    output_dir, item_id, 1
                )
                try:
                    with open(code_save, "w", encoding="utf-8") as fo:
                        fo.write(edited_code)
                except OSError as e:
                    print(f"[Batch][Failed] ID {item_id}: code write error: {e}")
                    fail += 1
                    continue

                # Code is on disk → success regardless of what happens next.
                success += 1

                if thoughts:
                    try:
                        with open(txt_save, "w", encoding="utf-8") as fo:
                            fo.write("\n\n".join(thoughts))
                    except OSError as e:
                        print(f"[Batch] ID {item_id}: thoughts write error: {e}")

                # Apply (best-effort). Failure is logged but never flips success.
                ok, err = apply_edit_and_save(
                    master_bytes=master_bytes,
                    slide_idx_0based=slide_idx,
                    edit_code=edited_code,
                    output_pptx_path=pptx_save,
                    target_w_emu=new_w_emu if target_ar else None,
                    target_h_emu=new_h_emu if target_ar else None,
                )
                if ok:
                    try:
                        with open(render_save, "w", encoding="utf-8") as fo:
                            json.dump({"width_emu": new_w_emu, "height_emu": new_h_emu}, fo)
                    except OSError as e:
                        print(f"[Batch] ID {item_id}: render meta write error: {e}")
                else:
                    print(f"[Batch] ID {item_id}: apply_edit failed ({err}); code preserved")
                    try:
                        if os.path.exists(pptx_save):
                            os.remove(pptx_save)
                    except OSError:
                        pass

    missing = set(prompt_to_item) - seen
    if missing:
        missing_ids = sorted({prompt_to_item[k][0] for k in missing})
        print(f"[Batch] {len(missing)} items had no row in predictions.jsonl: {missing_ids[:10]}...")
        fail += len(missing)

    return success + cached_success, fail


# ---------------------------------------------------------------------------
# Apply phase: code -> pptx for every item that has a saved .edit.py
# ---------------------------------------------------------------------------

def apply_all(data_list, output_dir, master_pptx, id_mapping):
    """Run apply_edit_and_save for every item with code but missing pptx.

    Independent of the LLM phase: lets us regenerate pptx for items whose
    code was produced in a previous run but never converted (or whose pptx
    was deleted). render_all consumes whatever pptx exists after this.

    Returns (success, fail, skipped) where:
      - success: items whose pptx was produced (or already present) by this call
      - fail:    items where apply_edit_and_save returned an error
      - skipped: items skipped because they have no code on disk yet
    """
    if not os.path.isdir(output_dir):
        return 0, 0, 0

    master_bytes = _read_master_bytes(master_pptx)
    master_prs = Presentation(io.BytesIO(master_bytes))
    n_slides = len(list(master_prs.slides))
    slide_w = master_prs.slide_width
    slide_h = master_prs.slide_height

    success = 0
    fail = 0
    skipped = 0

    for item in tqdm(data_list, desc="Applying code"):
        item_id = str(item.get("id"))
        target_ar = item.get("target_aspect_ratio")

        try:
            new_id_int = int(item_id)
        except ValueError:
            continue
        old_id = id_mapping.get(new_id_int)
        if old_id is None:
            continue
        slide_idx = old_id - 1
        if not (0 <= slide_idx < n_slides):
            continue

        new_w_emu, new_h_emu = (
            compute_new_dims_emu(slide_w, slide_h, target_ar) if target_ar else (slide_w, slide_h)
        )

        for var_idx in range(1, NUM_VARIANTS + 1):
            pptx_save, _, _, render_save, code_save = edited_paths(
                output_dir, item_id, var_idx
            )

            if not os.path.exists(code_save):
                skipped += 1
                continue

            if os.path.exists(pptx_save):
                # Already applied; don't redo.
                success += 1
                continue

            try:
                with open(code_save, "r", encoding="utf-8") as f:
                    edit_code = f.read()
            except OSError as e:
                tqdm.write(f"[Apply] ID {item_id} (Variant {var_idx}): cannot read code: {e}")
                fail += 1
                continue

            if not edit_code.strip():
                tqdm.write(f"[Apply] ID {item_id} (Variant {var_idx}): empty code file")
                fail += 1
                continue

            ok, err = apply_edit_and_save(
                master_bytes=master_bytes,
                slide_idx_0based=slide_idx,
                edit_code=edit_code,
                output_pptx_path=pptx_save,
                target_w_emu=new_w_emu if target_ar else None,
                target_h_emu=new_h_emu if target_ar else None,
            )
            if ok:
                with open(render_save, "w", encoding="utf-8") as f:
                    json.dump({"width_emu": new_w_emu, "height_emu": new_h_emu}, f)
                success += 1
            else:
                tqdm.write(f"[Apply] ID {item_id} (Variant {var_idx}): {err}")
                try:
                    if os.path.exists(pptx_save):
                        os.remove(pptx_save)
                except OSError:
                    pass
                fail += 1

    return success, fail, skipped


# ---------------------------------------------------------------------------
# Rendering: PPTX -> PDF -> PNG via LibreOffice (headless) + pdf2image
# ---------------------------------------------------------------------------

def _render_one(pptx_path_str, png_path_str, dpi, max_attempts=3):
    pptx_path = Path(pptx_path_str)
    png_path = Path(png_path_str)
    out_dir = png_path.parent
    pdf_path = out_dir / (pptx_path.stem + ".pdf")

    last_err = None
    for attempt in range(1, max_attempts + 1):
        try:
            # Per-process LibreOffice profile so parallel workers don't clash.
            profile_dir = Path("/tmp") / f"lo_profile_{os.getpid()}_{attempt}"
            profile_uri = profile_dir.as_uri()
            cmd = [
                "soffice", "--headless",
                f"-env:UserInstallation={profile_uri}",
                "--convert-to", "pdf",
                "--outdir", str(out_dir),
                str(pptx_path),
            ]
            subprocess.run(cmd, check=True, timeout=180, capture_output=True)

            if not pdf_path.exists():
                raise FileNotFoundError(f"LibreOffice produced no PDF at {pdf_path}")

            from pdf2image import convert_from_path
            images = convert_from_path(str(pdf_path), dpi=dpi)
            if not images:
                raise RuntimeError("pdf2image returned no pages")
            images[0].save(str(png_path))

            try:
                pdf_path.unlink()
            except OSError:
                pass

            suffix = "" if attempt == 1 else f" [attempt {attempt}/{max_attempts}]"
            return f"[OK] {png_path.name} (dpi={dpi}){suffix}"
        except subprocess.CalledProcessError as e:
            stderr = (e.stderr or b"").decode("utf-8", errors="replace")[:400]
            last_err = f"soffice exit {e.returncode}: {stderr}"
        except Exception as e:
            last_err = e
        if attempt < max_attempts:
            time.sleep(1)

    return f"[ERR] {pptx_path.name}: {last_err} (after {max_attempts} attempts)"


def render_all(output_dir, num_workers=4, dpi=DEFAULT_RENDER_DPI, force=False):
    """Render every *_edited_*.pptx in output_dir to a sibling .png.

    Returns (with_png, rendered_now, failed):
      - with_png:     total pptx that have a png after this call
                      (= pre-existing png + rendered_now)
      - rendered_now: png files produced in this call
      - failed:       pptx that failed to render
    """
    out = Path(output_dir)
    if not out.is_dir():
        print(f"[Render] {output_dir} does not exist")
        return 0, 0, 0
    pptx_files = sorted(out.glob("*_edited_*.pptx"))
    if not pptx_files:
        print(f"[Render] No PPTX files in {output_dir}")
        return 0, 0, 0

    tasks = []
    for pf in pptx_files:
        png = pf.with_suffix(".png")
        if png.exists() and not force:
            continue
        tasks.append((str(pf), str(png), dpi))

    pre_existing = len(pptx_files) - len(tasks)

    if not tasks:
        print(
            f"[Render] {len(pptx_files)} pptx in {output_dir}; "
            f"all {pre_existing} already have png — nothing to render"
        )
        return pre_existing, 0, 0

    print(
        f"[Render] {len(pptx_files)} pptx in {output_dir}: "
        f"{pre_existing} already have png, rendering {len(tasks)} with {num_workers} workers"
    )
    rendered = 0
    fail = 0
    with ProcessPoolExecutor(max_workers=num_workers) as ex:
        futs = {ex.submit(_render_one, p, png, d): p for p, png, d in tasks}
        for fut in tqdm(as_completed(futs), total=len(futs), desc="Rendering"):
            try:
                msg = fut.result()
            except Exception as e:
                print(f"[Render] Future exception: {e}")
                fail += 1
                continue
            if msg.startswith("[OK]"):
                rendered += 1
            else:
                print(msg)
                fail += 1
    return pre_existing + rendered, rendered, fail


# ---------------------------------------------------------------------------
# W&B logging + main
# ---------------------------------------------------------------------------

def upload_to_wandb(
    args, per_op_plan, total_tasks, success, fail,
    total_code_files=0, total_pptx_files=0, total_png_files=0,
):
    run = wandb.init(
        project=args.wandb_project,
        name=args.wandb_run_name,
        config={
            "model_path": args.model_path,
            "input_file": args.input_file,
            "output_dir": args.output_dir,
            "master_pptx": args.master_pptx,
            "id_mapping_csv": args.id_mapping_csv,
            "operations": list(per_op_plan.keys()),
            "input_files": [p["input_file"] for p in per_op_plan.values()],
            "output_dirs": [p["output_dir"] for p in per_op_plan.values()],
            "limit": args.limit,
            "num_workers": args.num_workers,
            "total_tasks": total_tasks,
            "num_variants": NUM_VARIANTS,
            "approach": "code-based-pptx-edit",
            "input_mode": args.input_mode,
        },
    )
    # Primary success metric is code presence (this run). pptx/png are tracked
    # separately as downstream artifact counts on disk.
    run.log({
        "total": success + fail,
        "code_success": success,
        "code_failed": fail,
        "code_success_rate": success / (success + fail) if (success + fail) > 0 else 0,
        "code_files_on_disk": total_code_files,
        "pptx_files_on_disk": total_pptx_files,
        "png_files_on_disk": total_png_files,
    })
    artifact = wandb.Artifact(
        name="edited-infographics-ppt",
        type="dataset",
        description=(
            f"Code-based edited PPT infographics — code: {success} success / {fail} failed; "
            f"on disk: code={total_code_files}, pptx={total_pptx_files}, png={total_png_files}"
        ),
    )
    for op, plan in per_op_plan.items():
        outd = plan["output_dir"]
        if os.path.isdir(outd):
            artifact.add_dir(os.path.abspath(outd), name=op)
    run.log_artifact(artifact)
    run.finish()
    print(f"W&B run uploaded: {run.url}")


def parse_args():
    parser = argparse.ArgumentParser(
        description="Code-based PPTX infographic editing: LLM modifies python-pptx code, then we render to PNG."
    )
    parser.add_argument("--api_keys", type=str, nargs="+", default=None,
                        help="One or more Google API Keys (online mode only).")
    parser.add_argument("--model_path", type=str, default=DEFAULT_MODEL_PATH)
    parser.add_argument("--input_file", type=str, default="editing_prompts.jsonl",
                        help="Base path of the editing-prompt file. The actual file per operation is "
                             "<stem>.<operation><ext>.")
    parser.add_argument("--output_dir", type=str, default="edited_infographics_ppt",
                        help="Base output directory; per-op artifacts go to <output_dir>/<operation>/.")
    parser.add_argument("--master_pptx", type=str, required=True,
                        help="Path to the master multi-slide .pptx (e.g. data/ppt_infographics/v7/tempates_gallery_0430.pptx).")
    parser.add_argument("--id_mapping_csv", type=str, required=True,
                        help="Path to id_mapping.csv (new_id,old_id). old_id is the 1-based slide index in master.")
    parser.add_argument("--operation", type=str, default="add",
                        choices=["add", "delete", "swap_intra", "swap_inter", "text_expand", "aspect_ratio", "all"])
    parser.add_argument("--input_mode", type=str, default="code", choices=list(INPUT_MODES))
    parser.add_argument("--prompt_index", type=int, default=0)
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--no_skip", action="store_true")
    parser.add_argument("--dry_run", action="store_true")
    parser.add_argument("--num_workers", type=int, default=1)
    parser.add_argument("--use_batch", action="store_true",
                        help="Use Vertex AI Gemini Batch Prediction.")
    parser.add_argument("--gcp_project", type=str, default=None)
    parser.add_argument("--gcp_location", type=str, default="us-central1")
    parser.add_argument("--batch_bucket_uri", type=str, default=None)
    parser.add_argument("--batch_poll_interval", type=int, default=30)
    parser.add_argument("--num_render_workers", type=int, default=4,
                        help="Parallel LibreOffice workers in the rendering phase. Each gets its own LO profile.")
    parser.add_argument("--render_dpi", type=int, default=DEFAULT_RENDER_DPI,
                        help="DPI for PDF-to-PNG conversion.")
    parser.add_argument("--skip_apply", action="store_true",
                        help="Skip the apply phase (code → pptx). By default this runs after the edit "
                             "phase so any item with code on disk gets a pptx, even ones that already "
                             "had code from a previous run.")
    parser.add_argument("--skip_render", action="store_true")
    parser.add_argument("--render_only", action="store_true")
    parser.add_argument("--force_render", action="store_true")
    parser.add_argument("--wandb_project", type=str, default="infographic-editing-ppt")
    parser.add_argument("--wandb_run_name", type=str, default=None)
    parser.add_argument("--no_wandb", action="store_true")
    return parser.parse_args()


def main():
    args = parse_args()

    if getattr(args, "use_batch", False):
        raise SystemExit(
            "--use_batch is a Vertex AI feature and is not available on the "
            "openai-proxy branch. Drop --use_batch and raise --num_workers "
            "instead; requests then go to the proxy concurrently."
        )

    if not os.path.exists(args.master_pptx):
        print(f"Error: master pptx not found: {args.master_pptx}")
        return
    if not os.path.exists(args.id_mapping_csv):
        print(f"Error: id_mapping csv not found: {args.id_mapping_csv}")
        return
    id_mapping = load_id_mapping(args.id_mapping_csv)
    print(f"Loaded id_mapping with {len(id_mapping)} entries from {args.id_mapping_csv}")

    operations = list(OPERATIONS) if args.operation == "all" else [args.operation]

    per_op_plan = {}
    for op in operations:
        input_file = input_file_for(args.input_file, op)
        output_dir = output_dir_for(args.output_dir, op)

        if args.render_only:
            per_op_plan[op] = {
                "input_file": input_file, "output_dir": output_dir,
                "data": [], "all_data": [],
            }
            continue

        if not os.path.exists(input_file):
            print(f"[{op}] Input file not found: {input_file} — skipping")
            continue

        data = load_prompts(input_file, prompt_index=args.prompt_index)
        for item in data:
            item["operation"] = op

        if args.limit is not None and args.limit > 0:
            data = data[: args.limit]
            print(f"[{op}] Limited to first {args.limit} tasks")

        # Keep the full (post-limit, pre-skip) list around so the apply phase
        # can revisit items that were filtered out below for already having code.
        all_data = list(data)

        # Code presence is the primary indicator: drop items that already have
        # an .edit.py for every variant — they don't need another LLM call.
        # pptx/png are downstream/consequential outputs and tracked separately.
        if not args.no_skip and os.path.isdir(output_dir):
            before = len(data)
            data = [d for d in data if not all_code_exist(output_dir, str(d.get("id")))]
            if before - len(data) > 0:
                print(f"[{op}] Skipping {before - len(data)} items with code already in {output_dir}")

        if os.path.isdir(output_dir) and data:
            no_code_ids = [str(d.get("id")) for d in data]
            print(f"[{op}] No code ({len(no_code_ids)} items, will call LLM): {no_code_ids}")

        per_op_plan[op] = {
            "input_file": input_file, "output_dir": output_dir,
            "data": data, "all_data": all_data,
        }

    total_tasks = sum(len(p["data"]) for p in per_op_plan.values())
    if total_tasks == 0 and not args.render_only and args.skip_render:
        print("Nothing to do.")
        return

    print("-" * 50)
    mode = "render-only" if args.render_only else ("batch (Vertex AI)" if args.use_batch else "online")
    print(f"Mode:        {mode}")
    print(f"Input Mode:  {args.input_mode}")
    print(f"Model:       {args.model_path}")
    print(f"Master PPTX: {args.master_pptx}")
    print(f"ID Mapping:  {args.id_mapping_csv}")
    print(f"Operations:  {', '.join(per_op_plan.keys())}")
    print(f"Total Tasks: {total_tasks}")
    for op, plan in per_op_plan.items():
        print(f"  - [{op}] {len(plan['data'])} tasks | in={plan['input_file']} | out={plan['output_dir']}/")
    print(f"Variants:    {NUM_VARIANTS}")
    if args.use_batch and not args.render_only:
        print(f"GCP Project:  {args.gcp_project}")
        print(f"GCP Location: {args.gcp_location}")
        print(f"Bucket URI:   {args.batch_bucket_uri}")
    elif not args.render_only:
        print(f"Num Workers: {args.num_workers}")
        print(f"Proxy:        {os.environ.get('OPENAI_BASE_URL', '(OPENAI_BASE_URL unset)')}")
    print(f"Render workers: {args.num_render_workers} | dpi={args.render_dpi} | skip_render={args.skip_render}")
    print("-" * 50)

    if args.dry_run:
        for op, plan in per_op_plan.items():
            print(f"\n[Dry Run] operation={op} — first 3 edit prompts:\n")
            for item in plan["data"][:3]:
                new_id = int(item.get("id"))
                old_id = id_mapping.get(new_id)
                print(f"--- ID {new_id} (master slide {old_id}) | {item.get('image_path')} ---")
                print(item.get("generated_edit_prompt", ""))
                print()
        return

    if not args.render_only:
        if args.use_batch:
            if not args.gcp_project or not args.batch_bucket_uri:
                print("Error: --use_batch requires --gcp_project and --batch_bucket_uri")
                return

    total_success = 0
    total_fail = 0
    for op, plan in per_op_plan.items():
        output_dir = plan["output_dir"]
        os.makedirs(output_dir, exist_ok=True)

        if not args.render_only and plan["data"]:
            print(f"\n=== [{op}] Editing PPTX for {len(plan['data'])} items → {output_dir}/ ===")
            if args.use_batch:
                s, f = process_edits_batch(
                    project_id=args.gcp_project,
                    location=args.gcp_location,
                    bucket_uri=args.batch_bucket_uri,
                    model_path=args.model_path,
                    data_list=plan["data"],
                    output_dir=output_dir,
                    master_pptx=args.master_pptx,
                    id_mapping=id_mapping,
                    poll_interval=args.batch_poll_interval,
                    input_mode=args.input_mode,
                )
            elif args.num_workers and args.num_workers > 1:
                s, f = run_parallel(
                    args.num_workers, args.api_keys, args.model_path,
                    plan["data"], output_dir, args.master_pptx, id_mapping,
                    input_mode=args.input_mode,
                )
            else:
                clients = [None]  # unused: llm_client holds the proxy client
                s, f = process_edits(
                    clients=clients,
                    model_path=args.model_path,
                    data_list=plan["data"],
                    output_dir=output_dir,
                    master_pptx=args.master_pptx,
                    id_mapping=id_mapping,
                    worker_id=0,
                    input_mode=args.input_mode,
                )
            print(f"[{op}] Edit phase (code obtained): {s} success / {f} failed")
        else:
            s, f = 0, 0

        # Apply phase: ensure every item with code on disk also has a pptx.
        # Catches items the edit phase skipped (they already had code) but
        # whose pptx is missing (e.g. apply failed last time, or pptx was
        # deleted). Skipped in render-only mode and via --skip_apply.
        if not args.render_only and not args.skip_apply and plan.get("all_data"):
            print(f"\n=== [{op}] Applying code → pptx ===")
            apply_s, apply_f, apply_skipped = apply_all(
                plan["all_data"], output_dir, args.master_pptx, id_mapping,
            )
            print(
                f"[{op}] Apply phase: {apply_s} ok / {apply_f} failed "
                f"({apply_skipped} variants had no code yet)"
            )

        if not args.skip_render:
            print(f"\n=== [{op}] Rendering PPTX to PNG ===")
            png_total, png_rendered, png_failed = render_all(
                output_dir,
                num_workers=args.num_render_workers,
                dpi=args.render_dpi,
                force=args.force_render,
            )
            print(
                f"[{op}] Render phase: {png_total} pptx have png "
                f"({png_rendered} rendered now, {png_failed} failed)"
            )

        total_success += s
        total_fail += f

    # Primary success metric: number of items with code (this run only).
    print(f"\nCode success: {total_success} / {total_tasks}")
    print(f"Code failed:  {total_fail} / {total_tasks}")

    # Separate metrics counted from disk: pptx and png are downstream of code.
    print("\nArtifacts on disk per op:")
    total_code_files = 0
    total_pptx_files = 0
    total_png_files = 0
    for op, plan in per_op_plan.items():
        outd = plan["output_dir"]
        if not os.path.isdir(outd):
            print(f"  - [{op}] (missing) {outd}/")
            continue
        out_path = Path(outd)
        code_n = len(list(out_path.glob("*_edited_*.edit.py")))
        pptx_n = len(list(out_path.glob("*_edited_*.pptx")))
        png_n = len(list(out_path.glob("*_edited_*.png")))
        total_code_files += code_n
        total_pptx_files += pptx_n
        total_png_files += png_n
        print(f"  - [{op}] code={code_n} pptx={pptx_n} png={png_n} | {outd}/")
    print(f"\nTotal code files: {total_code_files}")
    print(f"Total pptx files: {total_pptx_files}")
    print(f"Total png files (separate metric): {total_png_files}")

    if not args.no_wandb and not args.render_only:
        upload_to_wandb(
            args, per_op_plan, total_tasks, total_success, total_fail,
            total_code_files=total_code_files,
            total_pptx_files=total_pptx_files,
            total_png_files=total_png_files,
        )


if __name__ == "__main__":
    main()
